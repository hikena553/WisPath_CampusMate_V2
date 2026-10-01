from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query, Body
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import func
from urllib.parse import quote
from typing import Any
from collections import defaultdict
from pydantic import BaseModel

from app.core.database import get_db
from app.core.deps import require_role
from app.core.security import hash_password
from app.models.user import User, UserRole
from app.models.crisis import AIDialogSummary
from app.models.academic import Course, ClassGroup, Major, College, Semester
from app.models.knowledge import KnowledgeItem
from app.models.document import Document
from app.models.setting import SystemSetting
from app.models.conversation import Conversation, ConversationMessage
from app.core.crypto import encrypt_value, decrypt_value, is_encrypted
from app.schemas.admin import (
    KnowledgeItemCreate, KnowledgeItemUpdate, KnowledgeItemOut,
    DocumentOut, TeacherCreate, TeacherOut, StudentBriefOut, StudentUpdate, ImportResult,
)
from app.schemas.academic import CourseOut, CourseCreate, CourseImportResult
from app.services import knowledge_service
from app.services.import_export_service import export_users, import_users
from app.services.scoring import calc_radar_score

router = APIRouter(prefix="/api/admin", tags=["admin"])


# ========== 分页响应模型 ==========
class PaginatedResponse(BaseModel):
    """通用分页响应模型"""
    items: list[Any]
    total: int
    page: int
    page_size: int
    total_pages: int


# ========== 知识库管理 ==========

@router.get("/knowledge", response_model=PaginatedResponse, summary="获取知识库列表", description="分页获取知识库问答对列表，支持按分类和关键词筛选")
def list_knowledge(
    category: str | None = Query(None, description="分类筛选"),
    search: str | None = Query(None, description="搜索关键词"),
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(50, ge=1, le=200, description="每页条数"),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """获取知识库列表（分页）"""
    query = knowledge_service.get_all_knowledge_items(db, category, search)
    total = query.count()
    total_pages = (total + page_size - 1) // page_size
    start = (page - 1) * page_size
    items = query.offset(start).limit(page_size).all()

    return PaginatedResponse(
        items=[KnowledgeItemOut.model_validate(item) for item in items],
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@router.post("/knowledge", response_model=KnowledgeItemOut)
def create_knowledge(
    data: KnowledgeItemCreate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    return knowledge_service.create_knowledge_item(db, data.category, data.question, data.answer, data.tags)


@router.put("/knowledge/{item_id}", response_model=KnowledgeItemOut)
def update_knowledge(
    item_id: int,
    data: KnowledgeItemUpdate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    item = knowledge_service.update_knowledge_item(db, item_id, **data.model_dump(exclude_unset=True))
    if not item:
        raise HTTPException(status_code=404, detail="知识库条目不存在")
    return item


@router.delete("/knowledge/{item_id}")
def delete_knowledge(
    item_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    if not knowledge_service.delete_knowledge_item(db, item_id):
        raise HTTPException(status_code=404, detail="知识库条目不存在")
    return {"message": "删除成功"}


@router.post("/knowledge/upload")
async def upload_document(
    file: UploadFile = File(...),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    allowed_types = {"pdf", "docx", "txt"}
    filename = file.filename or "unnamed"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in allowed_types:
        raise HTTPException(status_code=400, detail="仅支持 PDF/DOCX/TXT 格式")

    import os
    from pathlib import Path
    upload_dir = Path(__file__).resolve().parent.parent.parent / "uploads" / "documents"
    upload_dir.mkdir(parents=True, exist_ok=True)

    # #6 路径清洗：只保留文件名部分，防止路径穿越
    safe_name = Path(filename).name
    file_path = upload_dir / safe_name

    # #7 大小限制：10MB
    MAX_UPLOAD_SIZE = 10 * 1024 * 1024
    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail="文件大小超过 10MB 限制")

    with open(file_path, "wb") as f:
        f.write(content)

    doc = knowledge_service.save_document(db, file.filename, ext, str(file_path), user.id)

    try:
        chunks = knowledge_service.parse_document(str(file_path), ext)
        if chunks:
            knowledge_service.save_chunks(db, doc.id, chunks)
        else:
            doc.status = "failed"
            db.commit()
            raise HTTPException(status_code=422, detail="文档解析失败：未能提取到有效内容")
    except HTTPException:
        raise
    except Exception as e:
        doc.status = "failed"
        db.commit()
        raise HTTPException(status_code=500, detail=f"文档解析失败: {str(e)}")

    return {"message": "上传成功", "document_id": doc.id, "chunk_count": len(chunks)}


@router.get("/knowledge/documents", response_model=list[DocumentOut])
def list_documents(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    docs = knowledge_service.get_all_documents(db)
    embedded_counts = knowledge_service.get_embedded_counts(db, [d.id for d in docs])
    return [DocumentOut(
        id=d.id,
        filename=d.filename,
        file_type=d.file_type,
        status=d.status,
        chunk_count=d.chunk_count,
        embedded_count=embedded_counts.get(d.id, 0),
        created_at=d.created_at.isoformat() if d.created_at else None,
    ) for d in docs]


@router.post("/knowledge/reindex")
def reindex_knowledge(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """为知识库文档分块重建向量索引（幂等：仅处理缺少向量的分块）"""
    result = knowledge_service.build_embeddings_for_missing(db)
    model = knowledge_service.get_embedding_model(db)
    return {
        "message": "向量索引重建完成" if not result["error"] else "向量索引重建失败",
        "model": model,
        **{k: v for k, v in result.items()},
    }


@router.get("/knowledge/index-status")
def knowledge_index_status(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """知识库检索模型与向量索引状态（分词模型/向量模型/索引覆盖率）"""
    return knowledge_service.get_index_status(db)


class EmbeddingConfigOut(BaseModel):
    model: str = ""
    base_url: str = ""
    api_key_set: bool = False
    configured: bool = False
    using_env_fallback: bool = False
    hint: str = ""


class EmbeddingConfigBody(BaseModel):
    model: str = ""          # 空串 = 使用默认 text-embedding-v3
    base_url: str = ""       # 空串 = 沿用环境变量 LLM_BASE_URL
    api_key: str = ""        # 明文，非空时加密存储


@router.get("/knowledge/embedding-config", response_model=EmbeddingConfigOut)
def embedding_get_config(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """读取向量模型配置（密钥仅返回是否已设置；显示是否回退环境变量）"""
    from app.core.config import settings

    model = _xqe_get(db, "embedding_model", "") or ""
    base = _xqe_get(db, "embedding_base_url", "") or ""
    api_key = _xqe_get(db, "embedding_api_key", "") or ""
    env_key = (settings.DASHSCOPE_API_KEY or settings.LLM_API_KEY or "").strip()
    env_base = (settings.LLM_BASE_URL or "").strip().rstrip("/")

    using_fallback = (not base or not api_key) and bool(env_key and env_base)
    effective_base = base or env_base
    effective_model = model or "text-embedding-v3"
    return EmbeddingConfigOut(
        model=effective_model,
        base_url=effective_base or "",
        api_key_set=bool(api_key or env_key),
        configured=bool((api_key or env_key) and (base or env_base)),
        using_env_fallback=using_fallback,
        hint=(
            "当前使用环境变量配置（DASHSCOPE_API_KEY / LLM_BASE_URL）"
            if using_fallback
            else ("请在下方填写模型名、接口地址与密钥" if not api_key else "配置已保存")
        ),
    )


@router.post("/knowledge/embedding-config", response_model=EmbeddingConfigOut)
def embedding_save_config(
    data: EmbeddingConfigBody,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """保存向量模型配置（模型名/接口地址/API 密钥，密钥加密存储）"""
    model = data.model.strip()
    if model:
        _xqe_set(db, "embedding_model", model)
    base = data.base_url.strip().rstrip("/")
    if base:
        if not base.startswith(("http://", "https://")):
            raise HTTPException(400, "接口地址需以 http:// 或 https:// 开头")
        _xqe_set(db, "embedding_base_url", base)
    if data.api_key.strip():
        _xqe_set(db, "embedding_api_key", encrypt_value(data.api_key.strip()))

    db.commit()
    return embedding_get_config(user=user, db=db)


@router.delete("/knowledge/documents/{doc_id}")
def delete_document(
    doc_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    if not knowledge_service.delete_document(db, doc_id):
        raise HTTPException(status_code=404, detail="文档不存在")
    return {"message": "删除成功"}


# ========== 教师管理 ==========

@router.post("/teachers", response_model=TeacherOut)
def create_teacher(
    data: TeacherCreate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    existing = db.query(User).filter(User.username == data.username).first()
    if existing:
        raise HTTPException(status_code=400, detail="工号已存在")

    teacher = User(
        username=data.username,
        name=data.name,
        role=UserRole.TEACHER,
        password_hash=hash_password("123456"),
        college=data.college,
        title=data.title,
        department=data.department,
        gender=data.gender,
        phone=data.phone,
    )
    db.add(teacher)
    db.commit()
    db.refresh(teacher)
    return TeacherOut(
        id=teacher.id, username=teacher.username, name=teacher.name,
        college=teacher.college, avatar=teacher.avatar,
        title=teacher.title, department=teacher.department,
        student_count=0,
    )


@router.get("/teachers", response_model=PaginatedResponse, summary="获取教师列表", description="分页获取教师列表，支持按姓名、工号、学院搜索")
def list_teachers(
    search: str | None = Query(None, description="搜索关键词"),
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(50, ge=1, le=200, description="每页条数"),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """获取教师列表（分页）"""
    # #9 子查询：一次性聚合每位教师的学生数量
    from sqlalchemy import func as sa_func
    student_count_sub = (
        db.query(User.tutor_id, sa_func.count(User.id).label("cnt"))
        .filter(User.role == UserRole.STUDENT, User.tutor_id.isnot(None))
        .group_by(User.tutor_id)
        .subquery()
    )

    query = (
        db.query(User, sa_func.coalesce(student_count_sub.c.cnt, 0).label("student_count"))
        .outerjoin(student_count_sub, User.id == student_count_sub.c.tutor_id)
        .filter(User.role == UserRole.TEACHER)
    )
    if search:
        from app.services.knowledge_service import _escape_like
        safe = _escape_like(search)
        like = f"%{safe}%"
        query = query.filter(User.name.like(like, escape="\\") | User.username.like(like, escape="\\") | User.college.like(like, escape="\\"))

    total = query.count()
    total_pages = (total + page_size - 1) // page_size
    start = (page - 1) * page_size
    rows = query.offset(start).limit(page_size).all()

    result = []
    for t, student_count in rows:
        result.append(TeacherOut(
            id=t.id,
            username=t.username,
            name=t.name,
            college=t.college,
            avatar=t.avatar,
            title=t.title,
            department=t.department,
            student_count=student_count,
        ))

    return PaginatedResponse(
        items=result,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@router.get("/teachers/{teacher_id}/students", response_model=list[StudentBriefOut])
def get_teacher_students(
    teacher_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    teacher = db.query(User).filter(User.id == teacher_id, User.role == UserRole.TEACHER).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="教师不存在")

    students = db.query(User).filter(User.role == UserRole.STUDENT, User.tutor_id == teacher_id).all()
    result = []
    for s in students:
        score = calc_radar_score(db, s.id, s)
        latest_crisis = db.query(AIDialogSummary).filter(
            AIDialogSummary.student_id == s.id
        ).order_by(AIDialogSummary.created_at.desc()).first()

        result.append(StudentBriefOut(
            id=s.id,
            username=s.username,
            name=s.name,
            college=s.college,
            avatar=s.avatar,
            score=score,
            crisis_level=latest_crisis.level.value if latest_crisis else None,
        ))
    return result


@router.delete("/teachers/batch")
def batch_delete_teachers(
    ids: list[int] = Body(..., embed=True),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """批量删除教师（注意：必须注册在 /teachers/{teacher_id} 之前，避免路由被吞）"""
    teachers = db.query(User).filter(User.id.in_(ids), User.role == UserRole.TEACHER).all()
    if not teachers:
        raise HTTPException(status_code=404, detail="未找到指定教师")

    deleted_names = []
    for t in teachers:
        db.query(User).filter(User.role == UserRole.STUDENT, User.tutor_id == t.id).update(
            {"tutor_id": None}
        )
        deleted_names.append(t.name)
        db.delete(t)

    db.commit()
    return {"message": f"已删除 {len(deleted_names)} 名教师：{', '.join(deleted_names)}"}


@router.delete("/teachers/{teacher_id}")
def delete_teacher(
    teacher_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    teacher = db.query(User).filter(User.id == teacher_id, User.role == UserRole.TEACHER).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="教师不存在")

    db.query(User).filter(User.role == UserRole.STUDENT, User.tutor_id == teacher_id).update(
        {"tutor_id": None}
    )
    db.delete(teacher)
    db.commit()
    return {"message": f"已删除教师 {teacher.name}，其名下学生的辅导员已清空"}


# ========== 学生管理 ==========

@router.get("/students", response_model=PaginatedResponse, summary="获取学生列表", description="分页获取学生列表，支持按学号、姓名、学院、班级筛选")
def list_students(
    search: str | None = Query(None, description="搜索学号或姓名"),
    college: str | None = Query(None, description="学院筛选"),
    class_name: str | None = Query(None, description="班级筛选"),
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(50, ge=1, le=200, description="每页条数"),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """获取学生列表（分页）"""
    query = db.query(User).filter(User.role == UserRole.STUDENT)
    if search:
        from app.services.knowledge_service import _escape_like
        safe = _escape_like(search)
        like = f"%{safe}%"
        query = query.filter(User.name.like(like, escape="\\") | User.username.like(like, escape="\\"))
    if college:
        from app.services.knowledge_service import _escape_like
        safe_college = _escape_like(college)
        query = query.filter(User.college.like(f"%{safe_college}%", escape="\\"))
    if class_name:
        query = query.filter(User.class_name == class_name)

    total = query.count()
    total_pages = (total + page_size - 1) // page_size
    start = (page - 1) * page_size
    students = query.offset(start).limit(page_size).all()

    # #10 批量查询：一次获取本页所有学生的最新危机记录
    student_ids = [s.id for s in students]
    latest_crisis_map: dict[int, str] = {}
    if student_ids:
        from sqlalchemy import func as sa_func
        crisis_sub = (
            db.query(
                AIDialogSummary.student_id,
                AIDialogSummary.level,
                sa_func.row_number().over(
                    partition_by=AIDialogSummary.student_id,
                    order_by=AIDialogSummary.created_at.desc()
                ).label("rn")
            )
            .filter(AIDialogSummary.student_id.in_(student_ids))
            .subquery()
        )
        crisis_rows = db.query(crisis_sub).filter(crisis_sub.c.rn == 1).all()
        latest_crisis_map = {r.student_id: r.level.value for r in crisis_rows}

    result = []
    for s in students:
        score = calc_radar_score(db, s.id, s)
        result.append(StudentBriefOut(
            id=s.id, username=s.username, name=s.name,
            college=s.college, class_name=s.class_name, avatar=s.avatar,
            score=score,
            crisis_level=latest_crisis_map.get(s.id),
        ))

    return PaginatedResponse(
        items=result,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@router.get("/students/stats")
def student_stats(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    total = db.query(User).filter(User.role == UserRole.STUDENT).count()

    college_rows = db.query(User.college, func.count(User.id)).filter(
        User.role == UserRole.STUDENT, User.college.isnot(None)
    ).group_by(User.college).all()
    college_stats = [{"college": c, "count": n} for c, n in college_rows]

    crisis_sub = db.query(
        AIDialogSummary.student_id,
        AIDialogSummary.level,
        func.row_number().over(
            partition_by=AIDialogSummary.student_id,
            order_by=AIDialogSummary.created_at.desc()
        ).label("rn")
    ).subquery()

    crisis_counts = db.query(crisis_sub.c.level, func.count(crisis_sub.c.student_id)).filter(
        crisis_sub.c.rn == 1
    ).group_by(crisis_sub.c.level).all()
    crisis_stats = [{"level": level.value, "count": n} for level, n in crisis_counts]

    no_crisis = total - sum(n for _, n in crisis_counts)
    if no_crisis > 0:
        crisis_stats.append({"level": "none", "count": no_crisis})

    return {"total": total, "college_stats": college_stats, "crisis_stats": crisis_stats}


@router.put("/students/{student_id}", response_model=StudentBriefOut)
def update_student(
    student_id: int,
    data: StudentUpdate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    student = db.query(User).filter(User.id == student_id, User.role == UserRole.STUDENT).first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")

    for key, val in data.model_dump(exclude_unset=True).items():
        if val is not None:
            setattr(student, key, val)

    db.commit()
    db.refresh(student)

    score = calc_radar_score(db, student.id, student)
    latest_crisis = db.query(AIDialogSummary).filter(
        AIDialogSummary.student_id == student.id
    ).order_by(AIDialogSummary.created_at.desc()).first()

    return StudentBriefOut(
        id=student.id, username=student.username, name=student.name,
        college=student.college, class_name=student.class_name, avatar=student.avatar,
        score=score,
        crisis_level=latest_crisis.level.value if latest_crisis else None,
    )


# ========== 密码管理 ==========

@router.post("/reset-password/{user_id}")
def reset_password(
    user_id: int,
    current_user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    target = db.query(User).filter(User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="用户不存在")

    target.password_hash = hash_password("123456")
    target.password_changed = False
    db.commit()
    return {"message": f"已将 {target.name} 的密码重置为 123456"}


# ========== 数据导入导出 ==========


# ========== 课程管理 ==========

@router.get("/export")
def export_data(
    role: str = Query(..., description="student 或 teacher"),
    current_user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    if role not in ("student", "teacher"):
        raise HTTPException(status_code=400, detail="role 必须是 student 或 teacher")

    user_role = UserRole.STUDENT if role == "student" else UserRole.TEACHER
    data = export_users(db, user_role)

    filename = "学生数据.xlsx" if role == "student" else "教师数据.xlsx"
    # 使用 RFC 5987 编码处理中文文件名
    encoded_filename = quote(filename)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/import", response_model=ImportResult)
async def import_data(
    role: str = Query(..., description="student 或 teacher"),
    file: UploadFile = File(...),
    current_user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    if role not in ("student", "teacher"):
        raise HTTPException(status_code=400, detail="role 必须是 student 或 teacher")

    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 Excel 文件")

    content = await file.read()
    user_role = UserRole.STUDENT if role == "student" else UserRole.TEACHER
    result = import_users(db, user_role, content)
    return ImportResult(**result)


# ========== 课程管理 ==========

@router.get("/courses", response_model=list[CourseOut])
def admin_list_courses(
    class_group_id: int | None = Query(None),
    semester: str | None = Query(None),
    college_id: int | None = Query(None),
    major_id: int | None = Query(None),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """管理员获取课程列表，支持多级筛选"""
    q = db.query(Course)
    if class_group_id:
        q = q.filter(Course.class_group_id == class_group_id)
    elif major_id:
        cg_ids = [cg.id for cg in db.query(ClassGroup).filter(ClassGroup.major_id == major_id).all()]
        q = q.filter(Course.class_group_id.in_(cg_ids))
    elif college_id:
        major_ids = [m.id for m in db.query(Major).filter(Major.college_id == college_id).all()]
        cg_ids = [cg.id for cg in db.query(ClassGroup).filter(ClassGroup.major_id.in_(major_ids)).all()]
        q = q.filter(Course.class_group_id.in_(cg_ids))
    if semester:
        q = q.filter(Course.semester == semester)
    return q.order_by(Course.day_of_week, Course.start_period).all()


@router.get("/courses/summary", summary="课程表总览统计", description="统计各学期课程表数量；以及指定学期范围内每个班级的课程编排情况（含未编排班级）")
def admin_courses_summary(
    semester: str | None = Query(None, description="学期，不传则返回全部学期的统计；班级维度的编排情况按该学期统计"),
    college_id: int | None = Query(None),
    major_id: int | None = Query(None),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """课程表总览：学期维度数量 + 班级维度编排情况（未编排班级也返回 course_count=0）"""
    # 1) 班级范围（outerjoin 专业/学院，便于前端直接展示名称）
    cg_q = (
        db.query(
            ClassGroup.id,
            ClassGroup.name,
            ClassGroup.grade,
            ClassGroup.student_count,
            Major.name.label("major_name"),
            College.name.label("college_name"),
        )
        .outerjoin(Major, ClassGroup.major_id == Major.id)
        .outerjoin(College, Major.college_id == College.id)
    )
    if college_id:
        cg_q = cg_q.filter(Major.college_id == college_id)
    elif major_id:
        cg_q = cg_q.filter(ClassGroup.major_id == major_id)

    schedules: list[dict] = []
    for row in cg_q.order_by(ClassGroup.grade.desc(), ClassGroup.name.asc()).all():
        schedules.append({
            "class_group_id": row.id,
            "class_name": row.name,
            "grade": row.grade,
            "student_count": row.student_count,
            "major_name": row.major_name or "",
            "college_name": row.college_name or "",
            "course_count": 0,
            "filled_slots": 0,
            "total_slots": 35,
        })
    scope_cg_ids = [s["class_group_id"] for s in schedules]

    # 2) 学期维度统计（受学院/专业筛选影响）
    semesters: list[dict] = []
    if scope_cg_ids:
        sem_rows = (
            db.query(
                Course.semester,
                func.count(Course.id),
                func.count(func.distinct(Course.class_group_id)),
            )
            .filter(Course.class_group_id.in_(scope_cg_ids))
            .group_by(Course.semester)
            .all()
        )
        for value, course_count, schedule_count in sem_rows:
            if not value:
                continue
            label = str(value)
            try:
                parts = label.split("-")
                if len(parts) == 3:
                    label = f"{parts[0]}-{parts[1]} 第{'一' if parts[2] == '1' else '二'}学期"
            except (IndexError, ValueError):
                pass
            semesters.append({
                "value": str(value),
                "label": label,
                "schedule_count": schedule_count,
                "course_count": course_count,
            })
        def _sem_sort_key(s: dict) -> tuple:
            try:
                parts = s["value"].split("-")
                return (-int(parts[0]), -int(parts[2]))
            except (IndexError, ValueError):
                return (0, 0)

        # 合并手工配置的学期（无课程数据也在课程表中展示，自定义名称优先）
        sem_by_value = {s["value"]: s for s in semesters}
        for sem in db.query(Semester).all():
            if sem.value in sem_by_value:
                sem_by_value[sem.value]["label"] = sem.label
            else:
                sem_by_value[sem.value] = {
                    "value": sem.value,
                    "label": sem.label,
                    "schedule_count": 0,
                    "course_count": 0,
                }
        semesters = sorted(sem_by_value.values(), key=_sem_sort_key)

    # 3) 班级维度：指定学期内的课程数与占用节次（7天 × 5个课次行 = 35 槽位）
    total_courses = 0
    if semester and scope_cg_ids:
        cg_map = {s["class_group_id"]: s for s in schedules}
        occupied: dict[int, set] = defaultdict(set)
        for c in db.query(Course).filter(
            Course.class_group_id.in_(scope_cg_ids),
            Course.semester == semester,
        ).all():
            cg = cg_map.get(c.class_group_id)
            if cg is None:
                continue
            cg["course_count"] += 1
            total_courses += 1
            for p in (1, 3, 5, 7, 9):
                if c.start_period <= p <= c.end_period:
                    occupied[c.class_group_id].add((c.day_of_week, p))
        for cg_id, rows in occupied.items():
            cg_map[cg_id]["filled_slots"] = len(rows)
    else:
        total_courses = sum(s["course_count"] for s in semesters)

    return {
        "semesters": semesters,
        "schedules": schedules,
        "total_courses": total_courses,
    }


@router.post("/courses", response_model=CourseOut)
def admin_create_course(
    data: CourseCreate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """管理员新增课程"""
    cg = db.query(ClassGroup).get(data.class_group_id)
    if not cg:
        raise HTTPException(400, "班级不存在")
    # 检测时间冲突
    conflict = db.query(Course).filter(
        Course.class_group_id == data.class_group_id,
        Course.semester == data.semester,
        Course.day_of_week == data.day_of_week,
        Course.start_period <= data.end_period,
        Course.end_period >= data.start_period,
    ).first()
    if conflict:
        raise HTTPException(400, f"时间冲突：与课程「{conflict.name}」(第{conflict.start_period}-{conflict.end_period}节)重叠")
    obj = Course(**data.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/courses/{course_id}", response_model=CourseOut)
def admin_update_course(
    course_id: int,
    data: CourseCreate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """管理员修改课程"""
    obj = db.query(Course).get(course_id)
    if not obj:
        raise HTTPException(404, "课程不存在")
    for k, v in data.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@router.delete("/courses/batch")
def admin_batch_delete_courses(
    ids: list[int] = Body(..., embed=True),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """管理员批量删除课程（注意：必须注册在 /courses/{course_id} 之前，避免路由被吞）"""
    count = db.query(Course).filter(Course.id.in_(ids)).delete(synchronize_session=False)
    db.commit()
    return {"ok": True, "deleted": count}


@router.delete("/courses/{course_id}")
def admin_delete_course(
    course_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """管理员删除课程"""
    obj = db.query(Course).get(course_id)
    if not obj:
        raise HTTPException(404, "课程不存在")
    db.delete(obj)
    db.commit()
    return {"ok": True}


@router.get("/semesters")
def admin_list_semesters(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """获取最近12个学期列表，从当前学期往前推算"""
    from datetime import date
    today = date.today()
    year = today.year
    month = today.month

    # 当前学期: 3-8月为春季学期(2), 9-2月为秋季学期(1)
    if 3 <= month <= 8:
        current = (year - 1, year, 2)
    else:
        current = (year, year + 1, 1)

    semesters: list[dict] = []
    y_start, y_end, term = current
    for i in range(12):
        semesters.append({
            "value": f"{y_start}-{y_end}-{term}",
            "label": f"{y_start}-{y_end} 第{'一' if term == 1 else '二'}学期",
        })
        if term == 1:
            y_start -= 1
            y_end -= 1
            term = 2
        else:
            term = 1

    # 同时合并数据库中已有的学期
    db_semesters = [r[0] for r in db.query(Course.semester).distinct().order_by(Course.semester.desc()).all()]
    existing_values = {s["value"] for s in semesters}
    for s in db_semesters:
        if s and s not in existing_values:
            semesters.append({"value": s, "label": s})
            existing_values.add(s)

    # 按学期倒序排列
    def sort_key(s: dict) -> tuple:
        try:
            parts = s["value"].split("-")
            return (-int(parts[0]), -int(parts[2]))
        except (IndexError, ValueError):
            return (0, 0)

    semesters.sort(key=sort_key)

    # 补全 label
    for s in semesters:
        if "第" not in s["label"]:
            parts = s["value"].split("-")
            if len(parts) == 3:
                s["label"] = f"{parts[0]}-{parts[1]} 第{'一' if parts[2] == '1' else '二'}学期"

    return semesters


class SemesterCreate(BaseModel):
    """新增学期配置"""
    year: str = Body(..., description="学年，如 2026-2027")
    term: int = Body(..., description="学期：1 第一学期，2 第二学期")
    label: str | None = Body(None, description="自定义显示名称，不传则按学年学期自动生成")


class SemesterUpdate(BaseModel):
    """编辑学期配置（仅允许修改显示名称，value 与课程数据强关联不可改）"""
    label: str = Body(...)


@router.get("/semesters/managed", summary="学期配置列表", description="返回手工配置的学期 + 数据库中已有课程的学期，附课程数量")
def admin_list_managed_semesters(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """学期管理列表：配置表学期 ∪ 课程表学期（去重），带各自课程数量"""
    import re
    from collections import defaultdict as _dd

    # 1) 配置表学期
    rows: list[dict] = []
    for sem in db.query(Semester).order_by(Semester.id.asc()).all():
        rows.append({
            "id": sem.id,
            "value": sem.value,
            "label": sem.label,
            "managed": True,
        })

    # 2) 课程表已有学期（补全 label、去重）
    def _fmt_label(v: str) -> str:
        parts = v.split("-")
        if len(parts) == 3 and re.fullmatch(r"\d{4}-\d{4}-\d", v):
            return f"{parts[0]}-{parts[1]} 第{'一' if parts[2] == '1' else '二'}学期"
        return v

    existing_values = {r["value"] for r in rows}
    for (v,) in db.query(Course.semester).distinct().all():
        if v and v not in existing_values:
            rows.append({"id": None, "value": v, "label": _fmt_label(v), "managed": False})
            existing_values.add(v)

    # 3) 课程数量统计
    counts: dict[str, int] = _dd(int)
    for (v, c) in db.query(Course.semester, func.count(Course.id)).group_by(Course.semester).all():
        if v:
            counts[v] = c
    for r in rows:
        r["course_count"] = counts.get(r["value"], 0)

    # 4) 排序：学年倒序、学期倒序
    def _k(r: dict) -> tuple:
        m = re.fullmatch(r"(\d{4})-(\d{4})-(\d)", r["value"] or "")
        return (-int(m.group(1)), -int(m.group(3))) if m else (0, 0)

    rows.sort(key=_k)
    return rows


@router.post("/semesters", summary="新增学期配置", description="新增一个学期（无课程数据时也会在课程表总览中展示）")
def admin_create_semester(
    data: SemesterCreate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    import re
    if not re.fullmatch(r"\d{4}-\d{4}", data.year.strip()):
        raise HTTPException(400, detail="学年格式不正确，示例：2026-2027")
    if data.term not in (1, 2):
        raise HTTPException(400, detail="学期只能选择第一学期或第二学期")

    value = f"{data.year.strip()}-{data.term}"

    # 与配置表 / 课程表学期均不能重复
    existed = db.query(Semester).filter(Semester.value == value).first()
    if existed:
        raise HTTPException(400, detail=f"学期 {value} 已存在")
    if value in {r[0] for r in db.query(Course.semester).distinct().all() if r[0]}:
        raise HTTPException(400, detail=f"学期 {value} 已存在于课程数据中")

    label = (data.label or "").strip()
    if not label:
        label = f"{data.year.strip()} 第{'一' if data.term == 1 else '二'}学期"

    sem = Semester(value=value, label=label)
    db.add(sem)
    db.commit()
    db.refresh(sem)
    return {"id": sem.id, "value": sem.value, "label": sem.label, "course_count": 0}


@router.put("/semesters/{semester_id}", summary="编辑学期配置", description="仅允许修改显示名称（label），value 与课程数据强关联不可改")
def admin_update_semester(
    semester_id: int,
    data: SemesterUpdate,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    sem = db.get(Semester, semester_id)
    if not sem:
        raise HTTPException(404, detail="学期配置不存在")
    label = data.label.strip()
    if not label:
        raise HTTPException(400, detail="显示名称不能为空")
    sem.label = label
    db.commit()
    db.refresh(sem)
    return {"id": sem.id, "value": sem.value, "label": sem.label}


@router.delete("/semesters/{semester_id}", summary="删除学期配置", description="删除手工配置的学期；若该学期已有课程数据则拒绝删除")
def admin_delete_semester(
    semester_id: int,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    sem = db.get(Semester, semester_id)
    if not sem:
        raise HTTPException(404, detail="学期配置不存在")
    course_count = db.query(func.count(Course.id)).filter(Course.semester == sem.value).scalar() or 0
    if course_count > 0:
        raise HTTPException(400, detail=f"该学期已有 {course_count} 门课程，请先删除课程后再删除学期")
    db.delete(sem)
    db.commit()
    return {"ok": True}


@router.post("/courses/import", response_model=CourseImportResult)
async def admin_import_courses(
    file: UploadFile = File(...),
    college_id: int = Query(...),
    semester: str = Query(...),
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """导入课程表 Excel，自动匹配学院下的班级"""
    import io
    import openpyxl

    college = db.query(College).get(college_id)
    if not college:
        raise HTTPException(400, "学院不存在")

    # 加载该学院下所有班级
    major_ids = [m.id for m in db.query(Major).filter(Major.college_id == college_id).all()]
    class_groups = db.query(ClassGroup).filter(ClassGroup.major_id.in_(major_ids)).all() if major_ids else []
    cg_map: dict[str, int] = {}  # 班级名 -> id
    for cg in class_groups:
        cg_map[cg.name] = cg.id
        # 也支持简称匹配，如去掉学院前缀
        short = cg.name
        for m in db.query(Major).filter(Major.id == cg.major_id).all():
            short = short.replace(m.name, "").strip()
        if short and short != cg.name:
            cg_map[short] = cg.id

    result = CourseImportResult()
    result.matched_classes = list(cg_map.keys())

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        ws = wb.active

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        if not rows:
            raise HTTPException(400, "Excel 文件为空（至少需要表头行和一行数据）")

        headers = [str(h).strip() if h else "" for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

        def find_col(*keys: str) -> int:
            for i, h in enumerate(headers):
                hl = h.lower().replace(" ", "").replace("_", "")
                for k in keys:
                    if k.lower().replace(" ", "").replace("_", "") == hl:
                        return i
            return -1

        idx_class = find_col("班级名称", "班级", "className", "class_name", "classname")
        idx_name = find_col("课程名称", "课程", "courseName", "course_name", "coursename", "name")
        idx_teacher = find_col("授课教师", "教师", "teacher")
        idx_location = find_col("上课地点", "教室", "地点", "location", "classroom")
        idx_day = find_col("星期", "day_of_week", "day")
        idx_start = find_col("开始节次", "start_period", "start")
        idx_end = find_col("结束节次", "end_period", "end")
        idx_ws = find_col("开始周", "week_start", "weekstart")
        idx_we = find_col("结束周", "week_end", "weekend")
        idx_credit = find_col("学分", "credit")

        if idx_class < 0 or idx_name < 0:
            raise HTTPException(400, "Excel 缺少必要列：班级名称、课程名称")

        for row_idx, row in enumerate(rows, start=2):
            try:
                class_name = str(row[idx_class]).strip() if idx_class < len(row) and row[idx_class] else ""
                course_name = str(row[idx_name]).strip() if idx_name < len(row) and row[idx_name] else ""

                if not class_name or not course_name:
                    result.errors.append({"row": row_idx, "msg": "班级名称或课程名称为空"})
                    continue

                # 匹配班级
                cg_id = cg_map.get(class_name)
                if cg_id is None:
                    # 模糊匹配
                    for name, cid in cg_map.items():
                        if class_name in name or name in class_name:
                            cg_id = cid
                            break
                if cg_id is None:
                    if class_name not in result.unmatched_classes:
                        result.unmatched_classes.append(class_name)
                    result.errors.append({"row": row_idx, "msg": f"未匹配到班级「{class_name}」"})
                    continue

                teacher = str(row[idx_teacher]).strip() if idx_teacher >= 0 and idx_teacher < len(row) and row[idx_teacher] else "未知"
                location = str(row[idx_location]).strip() if idx_location >= 0 and idx_location < len(row) and row[idx_location] else "待定"

                day_val = row[idx_day] if idx_day >= 0 and idx_day < len(row) else None
                if isinstance(day_val, str):
                    day_map = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "1": 1, "2": 2, "3": 3, "4": 4, "5": 5}
                    day_val = day_map.get(day_val.strip().replace("周", "").replace("星期", ""), 1)
                day_of_week = int(day_val) if day_val is not None else 1
                if day_of_week < 1 or day_of_week > 5:
                    result.errors.append({"row": row_idx, "msg": f"无效的星期值「{day_val}」"})
                    continue

                start_period = int(row[idx_start]) if idx_start >= 0 and idx_start < len(row) and row[idx_start] is not None else 1
                end_period = int(row[idx_end]) if idx_end >= 0 and idx_end < len(row) and row[idx_end] is not None else 2
                week_start = int(row[idx_ws]) if idx_ws >= 0 and idx_ws < len(row) and row[idx_ws] is not None else 1
                week_end = int(row[idx_we]) if idx_we >= 0 and idx_we < len(row) and row[idx_we] is not None else 16
                credit = float(row[idx_credit]) if idx_credit >= 0 and idx_credit < len(row) and row[idx_credit] is not None else None

                result.total += 1

                # 检查重复
                dup = db.query(Course).filter(
                    Course.class_group_id == cg_id,
                    Course.semester == semester,
                    Course.name == course_name,
                    Course.day_of_week == day_of_week,
                    Course.start_period == start_period,
                ).first()
                if dup:
                    result.skipped += 1
                    result.errors.append({"row": row_idx, "msg": f"已存在相同课程「{course_name}」，跳过"})
                    continue

                obj = Course(
                    class_group_id=cg_id,
                    semester=semester,
                    name=course_name,
                    teacher=teacher,
                    location=location,
                    day_of_week=day_of_week,
                    start_period=start_period,
                    end_period=end_period,
                    week_start=week_start,
                    week_end=week_end,
                    credit=credit,
                )
                db.add(obj)
                result.created += 1
            except Exception as e:
                result.errors.append({"row": row_idx, "msg": str(e)})

        if result.created > 0:
            db.commit()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"解析 Excel 失败：{str(e)}")

    return result


# ========== 喜鹊儿（青果教务）课表同步 ==========

class XiqueConfigOut(BaseModel):
    configured: bool = False
    root_url: str = ""
    username: str = ""
    password_set: bool = False
    school_year: int = 0
    term: int = 0
    semester: str = ""
    hint: str = ""


class XiqueConfigBody(BaseModel):
    root_url: str = ""
    username: str = ""
    password: str = ""  # 明文，非空时保存 md5 后加密存储
    school_year: int = 0
    term: int = 0


class XiqueSyncBody(BaseModel):
    root_url: str | None = None
    username: str | None = None
    password: str | None = None
    school_year: int | None = None
    term: int | None = None


def _xqe_get(db: Session, key: str, default=None):
    s = db.query(SystemSetting).filter(SystemSetting.key == key).first()
    return s.value if s and s.value not in (None, "") else default


def _xqe_set(db: Session, key: str, value: str) -> None:
    s = db.query(SystemSetting).filter(SystemSetting.key == key).first()
    if s:
        s.value = value
    else:
        db.add(SystemSetting(key=key, value=value))


@router.get("/courses/xique/config", response_model=XiqueConfigOut)
def xique_get_config(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """读取喜鹊儿同步配置状态（密码仅返回是否已设置）"""
    root_url = _xqe_get(db, "xqe_root_url", "")
    username = _xqe_get(db, "xqe_username", "")
    password = _xqe_get(db, "xqe_password", "")
    try:
        school_year = int(_xqe_get(db, "xqe_school_year", "0") or 0)
    except ValueError:
        school_year = 0
    try:
        term = int(_xqe_get(db, "xqe_term", "0") or 0)
    except ValueError:
        term = 0

    semester = ""
    if school_year:
        try:
            from app.services.xqe_sync import xq_to_semester
            semester = xq_to_semester(school_year, term)
        except Exception:
            semester = ""

    return XiqueConfigOut(
        configured=bool(root_url and username and password),
        root_url=root_url,
        username=username,
        password_set=bool(password),
        school_year=school_year,
        term=term,
        semester=semester,
        hint="" if root_url else "请填写教务系统根地址，如 https://jwgl.mycc.edu.cn",
    )


@router.post("/courses/xique/config", response_model=XiqueConfigOut)
def xique_save_config(
    data: XiqueConfigBody,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """保存喜鹊儿同步配置（密码保存为 md5 后加密，不落库明文）"""
    from app.services.xqe_sync import md5_hex

    root_url = data.root_url.strip().rstrip("/")
    if not root_url.startswith(("http://", "https://")):
        raise HTTPException(400, "教务地址需以 http:// 或 https:// 开头")
    _xqe_set(db, "xqe_root_url", root_url)

    username = data.username.strip()
    if username:
        _xqe_set(db, "xqe_username", username)
    else:
        raise HTTPException(400, "学号不能为空")

    if data.password:
        # 只保存 md5(明文) 的密文，用于青果二次 MD5 登录协议
        _xqe_set(db, "xqe_password", encrypt_value(md5_hex(data.password)))

    if data.school_year >= 2000:
        _xqe_set(db, "xqe_school_year", str(data.school_year))
    if data.term in (0, 1):
        _xqe_set(db, "xqe_term", str(data.term))

    db.commit()
    # 返回最新配置状态
    return xique_get_config(user=user, db=db)


@router.post("/courses/sync/xique")
def xique_sync_courses(
    data: XiqueSyncBody,
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    """一键同步喜鹊儿课表到本地课程表（同步请求，约 3-10 秒）"""
    from app.services.xqe_sync import (
        sync_timetable, md5_hex,
        XqeLoginError, XqeNetworkError, XqeParseError,
    )

    root_url = (data.root_url or "").strip().rstrip("/") or _xqe_get(db, "xqe_root_url", "")
    username = (data.username or "").strip() or _xqe_get(db, "xqe_username", "")
    term = data.term if data.term in (0, 1) else None
    if term is None:
        try:
            term = int(_xqe_get(db, "xqe_term", "0") or 0)
        except ValueError:
            term = 0
    school_year = data.school_year or 0
    if not school_year:
        try:
            school_year = int(_xqe_get(db, "xqe_school_year", "0") or 0)
        except ValueError:
            school_year = 0

    if not root_url:
        raise HTTPException(400, "尚未配置教务地址，请先在「喜鹊儿同步」中保存配置")
    if not username:
        raise HTTPException(400, "尚未配置学号，请先在「喜鹊儿同步」中保存配置")
    if school_year < 2000:
        raise HTTPException(400, "学年（起始年份）未配置或无效，请先保存配置")

    # 密码优先级：请求体明文 -> 存储的 md5 密文
    if data.password:
        once_md5 = md5_hex(data.password)
    else:
        stored = _xqe_get(db, "xqe_password", "")
        if not stored:
            raise HTTPException(400, "尚未配置密码，请先在「喜鹊儿同步」中保存配置")
        decrypted = decrypt_value(stored) if is_encrypted(stored) else stored
        if "****" in decrypted:
            raise HTTPException(400, "密码已脱敏未保存，请重新输入密码")
        once_md5 = decrypted

    try:
        result = sync_timetable(
            db,
            root_url=root_url,
            username=username,
            once_md5_password=once_md5,
            school_year=school_year,
            term=term,
        )
    except XqeLoginError as e:
        raise HTTPException(400, f"喜鹊儿登录失败：{e}")
    except XqeNetworkError as e:
        raise HTTPException(502, f"教务网络异常：{e}")
    except XqeParseError as e:
        raise HTTPException(502, f"课表解析失败：{e}")

    db.commit()
    return {
        "success": True,
        "message": (
            f"同步完成：学生 {result.student_name}（{result.class_name}），"
            f"学期 {result.semester}，共 {result.total} 条课程，"
            f"新增 {result.created} 条，更新 {result.updated} 条，跳过 {result.skipped} 条"
        ),
        "semester": result.semester,
        "class_name": result.class_name,
        "total": result.total,
        "created": result.created,
        "updated": result.updated,
        "skipped": result.skipped,
    }


# ========== 仪表盘统计 ==========

@router.get("/dashboard")
def dashboard_stats(
    user: User = Depends(require_role(UserRole.ADMIN)),
    db: Session = Depends(get_db),
):
    teacher_count = db.query(User).filter(User.role == UserRole.TEACHER).count()
    student_count = db.query(User).filter(User.role == UserRole.STUDENT).count()
    college_count = db.query(College).count()
    knowledge_count = db.query(KnowledgeItem).count()
    document_count = db.query(Document).count()

    student_gender_rows = db.query(User.gender, func.count(User.id)).filter(
        User.role == UserRole.STUDENT, User.gender.isnot(None)
    ).group_by(User.gender).all()
    student_gender_stats = {g or "未知": c for g, c in student_gender_rows}

    teacher_gender_rows = db.query(User.gender, func.count(User.id)).filter(
        User.role == UserRole.TEACHER, User.gender.isnot(None)
    ).group_by(User.gender).all()
    teacher_gender_stats = {g or "未知": c for g, c in teacher_gender_rows}

    conversation_count = db.query(Conversation).count()
    message_count = db.query(ConversationMessage).count()

    college_rows = db.query(User.college, func.count(User.id)).filter(
        User.role == UserRole.STUDENT, User.college.isnot(None)
    ).group_by(User.college).all()
    college_stats = [{"college": c, "count": n} for c, n in college_rows]

    teacher_college_rows = db.query(User.college, func.count(User.id)).filter(
        User.role == UserRole.TEACHER, User.college.isnot(None)
    ).group_by(User.college).all()
    teacher_college_stats = [{"college": c, "count": n} for c, n in teacher_college_rows]

    crisis_sub = db.query(
        AIDialogSummary.student_id,
        AIDialogSummary.level,
        func.row_number().over(
            partition_by=AIDialogSummary.student_id,
            order_by=AIDialogSummary.created_at.desc()
        ).label("rn")
    ).subquery()

    crisis_counts = db.query(crisis_sub.c.level, func.count(crisis_sub.c.student_id)).filter(
        crisis_sub.c.rn == 1
    ).group_by(crisis_sub.c.level).all()
    crisis_stats = [{"level": level.value, "count": n} for level, n in crisis_counts]

    no_crisis = student_count - sum(n for _, n in crisis_counts)
    if no_crisis > 0:
        crisis_stats.append({"level": "none", "count": no_crisis})

    return {
        "teacher_count": teacher_count,
        "student_count": student_count,
        "college_count": college_count,
        "knowledge_count": knowledge_count,
        "document_count": document_count,
        "student_gender_stats": student_gender_stats,
        "teacher_gender_stats": teacher_gender_stats,
        "conversation_count": conversation_count,
        "message_count": message_count,
        "college_stats": college_stats,
        "teacher_college_stats": teacher_college_stats,
        "crisis_stats": crisis_stats,
    }
